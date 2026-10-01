"""
Professional DOCX exporter for ArchDraft.

Generates consistently formatted Word documents with:
  - Optional project logo in the header
  - Consistent font family and sizes throughout
  - Professional table formatting with borders and shading
  - Proper heading hierarchy with consistent styling
  - Table of Contents with page numbers
  - Centered figures with proper captions
  - Professional bullet and numbered lists
  - Consistent spacing and margins
"""

import re
import logging
from datetime import date
from pathlib import Path
from typing import Dict, Any, List, Optional, Tuple

logger = logging.getLogger(__name__)


class ProfessionalDocxExporter:
    """
    Professional DOCX exporter for generated HLDs.

    Usage:
        exporter = ProfessionalDocxExporter()
        path = exporter.export(parts, blueprint, generated_sections, output_path, logo_path)
    """

    # Document accent and heading colors
    ACCENT_COLOR = "C74634"
    H1_COLOR = "1F4E78"    # Dark blue for main headings
    H2_COLOR = "2E75B5"    # Medium blue
    H3_COLOR = "5B9BD5"    # Light blue
    H4_COLOR = "70AD47"    # Green for level 4

    # Font settings
    FONT_FAMILY = "Aptos"
    BODY_FONT_SIZE = 11  # pts
    H1_FONT_SIZE = 16
    H2_FONT_SIZE = 14
    H3_FONT_SIZE = 12
    H4_FONT_SIZE = 11
    CAPTION_FONT_SIZE = 10
    HEADER_FOOTER_SIZE = 9

    # Spacing
    PARAGRAPH_SPACING_BEFORE = 6  # pts
    PARAGRAPH_SPACING_AFTER = 6
    HEADING_SPACING_BEFORE = 12
    HEADING_SPACING_AFTER = 6
    LINE_SPACING = 1.15  # multiplier

    # Image settings
    IMAGE_WIDTH_INCHES = 5.5

    # Page margins (inches)
    TOP_MARGIN = 0.94
    BOTTOM_MARGIN = 0.94
    LEFT_MARGIN = 1.0
    RIGHT_MARGIN = 1.0

    def export(
        self,
        parts: Dict[str, str],
        blueprint: Dict[str, Any],
        generated_sections: Dict[str, Any],
        output_path: str,
        logo_path: Optional[str] = None,
    ) -> str:
        """
        Build and save the professionally formatted DOCX file.

        Args:
            parts: Dict with executive_summary, toc_content, list_of_figures,
                   list_of_tables, appendices_content
            blueprint: Blueprint metadata (customer_name, product, etc.)
            generated_sections: Dict of section_id -> GeneratedSection
            output_path: Absolute path for output .docx file
            logo_path: Optional path to a project logo image

        Returns:
            Absolute path string of saved file
        """
        try:
            from docx import Document
            from docx.shared import Inches, Pt, RGBColor
            from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
            from docx.enum.style import WD_STYLE_TYPE
        except ImportError:
            raise ImportError("python-docx required. Install: pip install python-docx")

        self._Inches = Inches
        self._Pt = Pt
        self._RGBColor = RGBColor
        self._WD_ALIGN = WD_ALIGN_PARAGRAPH
        self._WD_LINE_SPACING = WD_LINE_SPACING
        self._WD_STYLE_TYPE = WD_STYLE_TYPE

        doc = Document()

        # Apply document formatting
        self._set_page_margins(doc)
        self._configure_default_styles(doc)
        self._apply_heading_styles(doc)

        # Add cover page
        self._add_cover_page(doc, blueprint)

        # Add headers and footers
        self._add_headers_footers(doc, blueprint, logo_path)

        # Build image lookup
        image_lookup = self._build_image_lookup(generated_sections)

        # Front matter
        if parts.get("executive_summary"):
            self._parse_and_insert(doc, parts["executive_summary"], image_lookup)
            doc.add_page_break()

        # Table of Contents with page numbers
        if parts.get("toc_content"):
            self._insert_toc_with_page_numbers(doc, parts["toc_content"])

        if parts.get("list_of_figures"):
            self._parse_and_insert(doc, parts["list_of_figures"], image_lookup)

        if parts.get("list_of_tables"):
            self._parse_and_insert(doc, parts["list_of_tables"], image_lookup)

        if parts.get("toc_content") or parts.get("executive_summary"):
            doc.add_page_break()

        # Main sections
        sorted_sections = sorted(
            generated_sections.values(),
            key=lambda s: self._section_sort_key(s.get("section_number", "999"))
        )

        for section_dict in sorted_sections:
            self._insert_section(doc, section_dict, image_lookup)

        # Appendices
        if parts.get("appendices_content"):
            self._parse_and_insert(doc, parts["appendices_content"], image_lookup)

        doc.save(output_path)
        logger.info(f"Professional DOCX saved to {output_path}")
        return str(Path(output_path).absolute())

    # ===================================================================
    # Document Setup
    # ===================================================================

    def _set_page_margins(self, doc) -> None:
        """Apply document margins."""
        Inches = self._Inches
        for section in doc.sections:
            section.top_margin = Inches(self.TOP_MARGIN)
            section.bottom_margin = Inches(self.BOTTOM_MARGIN)
            section.left_margin = Inches(self.LEFT_MARGIN)
            section.right_margin = Inches(self.RIGHT_MARGIN)

    def _configure_default_styles(self, doc) -> None:
        """Configure Normal and other base styles for consistency."""
        Pt = self._Pt

        # Configure Normal style (body text)
        try:
            normal_style = doc.styles['Normal']
            normal_style.font.name = self.FONT_FAMILY
            normal_style.font.size = Pt(self.BODY_FONT_SIZE)

            # Set paragraph spacing
            normal_style.paragraph_format.space_before = Pt(self.PARAGRAPH_SPACING_BEFORE)
            normal_style.paragraph_format.space_after = Pt(self.PARAGRAPH_SPACING_AFTER)
            normal_style.paragraph_format.line_spacing_rule = self._WD_LINE_SPACING.MULTIPLE
            normal_style.paragraph_format.line_spacing = self.LINE_SPACING
        except KeyError:
            logger.warning("Normal style not found")

        # Configure List Bullet style
        try:
            bullet_style = doc.styles['List Bullet']
            bullet_style.font.name = self.FONT_FAMILY
            bullet_style.font.size = Pt(self.BODY_FONT_SIZE)
            bullet_style.paragraph_format.space_before = Pt(3)
            bullet_style.paragraph_format.space_after = Pt(3)
            bullet_style.paragraph_format.left_indent = self._Inches(0.25)
        except KeyError:
            logger.debug("List Bullet style not found")

        # Configure List Number style
        try:
            number_style = doc.styles['List Number']
            number_style.font.name = self.FONT_FAMILY
            number_style.font.size = Pt(self.BODY_FONT_SIZE)
            number_style.paragraph_format.space_before = Pt(3)
            number_style.paragraph_format.space_after = Pt(3)
            number_style.paragraph_format.left_indent = self._Inches(0.25)
        except KeyError:
            logger.debug("List Number style not found")

    def _apply_heading_styles(self, doc) -> None:
        """Apply consistent heading styles."""
        Pt = self._Pt

        heading_specs = [
            ("Heading 1", self.H1_FONT_SIZE, True, self.H1_COLOR, self.HEADING_SPACING_BEFORE, self.HEADING_SPACING_AFTER),
            ("Heading 2", self.H2_FONT_SIZE, True, self.H2_COLOR, self.HEADING_SPACING_BEFORE, self.HEADING_SPACING_AFTER),
            ("Heading 3", self.H3_FONT_SIZE, True, self.H3_COLOR, 10, 6),
            ("Heading 4", self.H4_FONT_SIZE, True, self.H4_COLOR, 10, 6),
        ]

        for style_name, size_pt, bold, color_hex, space_before, space_after in heading_specs:
            try:
                style = doc.styles[style_name]
                style.font.name = self.FONT_FAMILY
                style.font.size = Pt(size_pt)
                style.font.bold = bold
                style.font.italic = False
                style.font.color.rgb = self._hex_to_rgb(color_hex)

                # Spacing
                style.paragraph_format.space_before = Pt(space_before)
                style.paragraph_format.space_after = Pt(space_after)
                style.paragraph_format.keep_with_next = True  # Keep heading with following paragraph
            except KeyError:
                logger.debug(f"Style '{style_name}' not found")

    # ===================================================================
    # Cover Page
    # ===================================================================

    def _add_cover_page(self, doc, blueprint: Dict[str, Any]) -> None:
        """Add a project cover page."""
        Pt = self._Pt
        WD_ALIGN = self._WD_ALIGN

        customer = blueprint.get("customer_name", "Customer")
        product = blueprint.get("product", "Product")
        doc_type = blueprint.get("document_type", "High Level Design")
        project = blueprint.get("project_name", blueprint.get("title", f"{product} {doc_type}"))
        today_str = date.today().strftime("%B %d, %Y")

        # Vertical spacing
        for _ in range(8):
            doc.add_paragraph()

        # Document type (smaller, above title)
        type_para = doc.add_paragraph()
        type_para.alignment = WD_ALIGN.CENTER
        type_run = type_para.add_run(doc_type.upper())
        type_run.font.name = self.FONT_FAMILY
        type_run.font.size = Pt(14)
        type_run.font.color.rgb = self._hex_to_rgb(self.H2_COLOR)
        type_para.paragraph_format.space_after = Pt(6)

        # Main title
        title_para = doc.add_paragraph()
        title_para.alignment = WD_ALIGN.CENTER
        title_run = title_para.add_run(project)
        title_run.font.name = self.FONT_FAMILY
        title_run.font.size = Pt(24)
        title_run.font.bold = True
        title_run.font.color.rgb = self._hex_to_rgb(self.H1_COLOR)
        title_para.paragraph_format.space_after = Pt(24)

        # Horizontal line (using border)
        self._add_horizontal_line(doc, self.H1_COLOR)

        doc.add_paragraph()
        doc.add_paragraph()

        # Metadata table
        meta_items = [
            ("Customer:", customer),
            ("Product:", product),
            ("Document Type:", doc_type),
            ("Date:", today_str),
            ("Classification:", "Pending review"),
        ]

        for label, value in meta_items:
            para = doc.add_paragraph()
            para.alignment = WD_ALIGN.CENTER

            label_run = para.add_run(label + " ")
            label_run.font.name = self.FONT_FAMILY
            label_run.font.size = Pt(12)
            label_run.font.bold = True

            value_run = para.add_run(value)
            value_run.font.name = self.FONT_FAMILY
            value_run.font.size = Pt(12)

            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.space_after = Pt(3)

        doc.add_page_break()

    def _add_horizontal_line(self, doc, color_hex: str) -> None:
        """Add a horizontal decorative line."""
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            para = doc.add_paragraph()
            para.alignment = self._WD_ALIGN.CENTER
            para.paragraph_format.space_before = self._Pt(0)
            para.paragraph_format.space_after = self._Pt(0)

            pPr = para._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")

            bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single")
            bottom.set(qn("w:sz"), "24")  # Thicker line
            bottom.set(qn("w:space"), "1")
            bottom.set(qn("w:color"), color_hex)

            pBdr.append(bottom)
            pPr.append(pBdr)
        except Exception as e:
            logger.debug(f"Could not add horizontal line: {e}")

    # ===================================================================
    # Header / Footer
    # ===================================================================

    def _add_headers_footers(self, doc, blueprint: Dict[str, Any], logo_path: Optional[str] = None) -> None:
        """Add document headers and footers with an optional project logo."""
        Pt = self._Pt
        WD_ALIGN = self._WD_ALIGN

        product = blueprint.get("product", "")
        doc_type = blueprint.get("document_type", "HLD")
        customer = blueprint.get("customer_name", "")

        section = doc.sections[0]

        # ---- Header ----
        header = section.header
        if header.paragraphs:
            header_table = self._create_header_table(header, doc_type, product, customer, logo_path)

        # ---- Footer ----
        footer = section.footer
        if footer.paragraphs:
            fp = footer.paragraphs[0]
        else:
            fp = footer.add_paragraph()

        fp.clear()
        fp.alignment = WD_ALIGN.CENTER

        # Footer classification placeholder
        conf_run = fp.add_run("Classification pending review  |  Page ")
        conf_run.font.name = self.FONT_FAMILY
        conf_run.font.size = Pt(self.HEADER_FOOTER_SIZE)
        conf_run.font.color.rgb = self._RGBColor(0x60, 0x60, 0x60)

        # Page number field
        self._add_page_number_field(fp)

    def _create_header_table(self, header, doc_type: str, product: str, customer: str, logo_path: Optional[str]) -> None:
        """Create a header table with document info on left and logo on right."""
        try:
            from docx.table import Table
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            # Clear existing content
            for para in header.paragraphs:
                para.clear()

            # Create 1-row, 2-column table
            table = header.add_table(rows=1, cols=2, width=self._Inches(6.5))
            table.alignment = self._WD_ALIGN.LEFT

            # Remove borders
            self._remove_table_borders(table)

            # Left cell: Document title
            left_cell = table.rows[0].cells[0]
            left_para = left_cell.paragraphs[0]
            left_para.alignment = self._WD_ALIGN.LEFT

            title_text = f"{doc_type} – {product}"
            if customer:
                title_text += f" ({customer})"

            left_run = left_para.add_run(title_text)
            left_run.font.name = self.FONT_FAMILY
            left_run.font.size = self._Pt(self.HEADER_FOOTER_SIZE)
            left_run.font.color.rgb = self._RGBColor(0x60, 0x60, 0x60)

            # Right cell: optional project logo or product name
            right_cell = table.rows[0].cells[1]
            right_para = right_cell.paragraphs[0]
            right_para.alignment = self._WD_ALIGN.RIGHT

            if logo_path and Path(logo_path).exists():
                # Insert logo image
                try:
                    run = right_para.add_run()
                    run.add_picture(logo_path, height=self._Inches(0.3))
                except Exception as e:
                    logger.warning(f"Could not insert logo: {e}")
                    # Fallback to text
                    self._add_title_mark(right_para)
            else:
                self._add_title_mark(right_para)

            # Set column widths
            table.rows[0].cells[0].width = self._Inches(4.5)
            table.rows[0].cells[1].width = self._Inches(2.0)

        except Exception as e:
            logger.warning(f"Could not create header table: {e}")
            # Fallback to simple text header
            para = header.paragraphs[0]
            para.alignment = self._WD_ALIGN.RIGHT
            run = para.add_run(f"{doc_type} – {product}")
            run.font.size = self._Pt(self.HEADER_FOOTER_SIZE)

    def _add_title_mark(self, para) -> None:
        """Add a neutral product mark when no project logo is supplied."""
        mark = para.add_run("ARCHDRAFT")
        mark.font.name = self.FONT_FAMILY
        mark.font.size = self._Pt(9)
        mark.font.bold = True
        mark.font.color.rgb = self._hex_to_rgb(self.ACCENT_COLOR)

    def _remove_table_borders(self, table) -> None:
        """Remove all borders from a table."""
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            tbl = table._element
            tblPr = tbl.tblPr
            if tblPr is None:
                tblPr = OxmlElement('w:tblPr')
                tbl.insert(0, tblPr)

            tblBorders = OxmlElement('w:tblBorders')
            for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
                border = OxmlElement(f'w:{border_name}')
                border.set(qn('w:val'), 'none')
                border.set(qn('w:sz'), '0')
                border.set(qn('w:space'), '0')
                border.set(qn('w:color'), 'auto')
                tblBorders.append(border)

            tblPr.append(tblBorders)
        except Exception as e:
            logger.debug(f"Could not remove table borders: {e}")

    def _add_page_number_field(self, para) -> None:
        """Add automatic page numbering field."""
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            run = para.add_run()
            run.font.name = self.FONT_FAMILY
            run.font.size = self._Pt(self.HEADER_FOOTER_SIZE)
            run.font.color.rgb = self._RGBColor(0x60, 0x60, 0x60)

            # Add field characters
            fldChar1 = OxmlElement('w:fldChar')
            fldChar1.set(qn('w:fldCharType'), 'begin')

            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = 'PAGE'

            fldChar2 = OxmlElement('w:fldChar')
            fldChar2.set(qn('w:fldCharType'), 'end')

            run._r.append(fldChar1)
            run._r.append(instrText)
            run._r.append(fldChar2)

        except Exception as e:
            logger.debug(f"Page number field failed: {e}")
            para.add_run(" [page]").font.size = self._Pt(self.HEADER_FOOTER_SIZE)

    # ===================================================================
    # Table of Contents with Page Numbers
    # ===================================================================

    def _insert_toc_with_page_numbers(self, doc, toc_content: str) -> None:
        """Insert Table of Contents with proper page number support."""
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            # Add TOC heading
            toc_heading = doc.add_heading("Table of Contents", level=1)

            # Add TOC field
            para = doc.add_paragraph()
            run = para.add_run()

            fldChar1 = OxmlElement('w:fldChar')
            fldChar1.set(qn('w:fldCharType'), 'begin')

            instrText = OxmlElement('w:instrText')
            instrText.set(qn('xml:space'), 'preserve')
            instrText.text = 'TOC \\o "1-4" \\h \\z \\u'

            fldChar2 = OxmlElement('w:fldChar')
            fldChar2.set(qn('w:fldCharType'), 'separate')

            fldChar3 = OxmlElement('w:fldChar')
            fldChar3.set(qn('w:fldCharType'), 'end')

            run._r.append(fldChar1)
            run._r.append(instrText)
            run._r.append(fldChar2)
            run._r.append(OxmlElement('w:t'))
            run._r.append(fldChar3)

            para.paragraph_format.space_after = self._Pt(12)

            # Add instruction note
            note_para = doc.add_paragraph()
            note_run = note_para.add_run("(Right-click and select 'Update Field' to refresh page numbers)")
            note_run.font.name = self.FONT_FAMILY
            note_run.font.size = self._Pt(9)
            note_run.font.italic = True
            note_run.font.color.rgb = self._RGBColor(0x80, 0x80, 0x80)

            doc.add_page_break()

        except Exception as e:
            logger.warning(f"Could not create TOC field: {e}. Using simple TOC.")
            # Fallback: parse and insert as regular paragraphs
            self._parse_and_insert(doc, toc_content, {})
            doc.add_page_break()

    # ===================================================================
    # Image Lookup
    # ===================================================================

    def _build_image_lookup(self, generated_sections: Dict[str, Any]) -> Dict[str, Dict]:
        """Build image_id -> {storage_path, caption} lookup."""
        lookup: Dict[str, Dict] = {}
        for section in generated_sections.values():
            for img in section.get("images_inserted", []):
                img_id = img.get("image_id", "")
                if img_id:
                    lookup[img_id] = {
                        "storage_path": img.get("storage_path", ""),
                        "caption": img.get("caption", ""),
                    }
        return lookup

    # ===================================================================
    # Section Insertion
    # ===================================================================

    def _insert_section(self, doc, section_dict: Dict[str, Any], image_lookup: Dict[str, Dict]) -> None:
        """Insert one section with its content and images."""
        markdown_content = section_dict.get("markdown_content", "")
        images_inserted = section_dict.get("images_inserted", [])

        # Track inline image IDs
        inline_image_ids = set(re.findall(r'!\[[^\]]*\]\(([^)]+)\)', markdown_content))
        inline_image_ids |= set(re.findall(r'\{\{FIGURE:([^}]+)\}\}', markdown_content))

        self._parse_and_insert(doc, markdown_content, image_lookup)

        # Insert non-inline images
        for img in images_inserted:
            img_id = img.get("image_id", "")
            if img_id and img_id not in inline_image_ids:
                self._insert_image(doc, img_id, img.get("caption", ""), image_lookup)

    # ===================================================================
    # Markdown Parser
    # ===================================================================

    def _parse_and_insert(self, doc, markdown_str: str, image_lookup: Dict[str, Dict]) -> None:
        """Parse markdown and insert with professional formatting."""
        lines = markdown_str.split("\n")
        i = 0

        while i < len(lines):
            line = lines[i]
            stripped = line.strip()

            # Empty line
            if not stripped:
                i += 1
                continue

            # Code fence block
            if stripped.startswith("```"):
                i = self._process_code_fence(doc, lines, i)
                continue

            # Section separator
            if stripped == "---":
                # Add small spacing instead of page break
                doc.add_paragraph()
                i += 1
                continue

            # Markdown table
            if stripped.startswith("|") and stripped.endswith("|"):
                i = self._process_table(doc, lines, i)
                continue

            # Figure reference {{FIGURE:...}}
            brace_match = re.match(r'\{\{FIGURE:([^}]+)\}\}', stripped)
            if brace_match:
                image_id = brace_match.group(1).strip()
                entry = image_lookup.get(image_id, {})
                caption = entry.get("caption", "")
                self._insert_image(doc, image_id, caption, image_lookup)
                i += 1
                continue

            # Image reference ![caption](image_id)
            img_match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
            if img_match:
                caption = img_match.group(1)
                image_id = img_match.group(2)
                self._insert_image(doc, image_id, caption, image_lookup)
                i += 1
                continue

            # Headings
            if stripped.startswith("#### "):
                self._add_heading(doc, stripped[5:], 4)
                i += 1
                continue
            elif stripped.startswith("### "):
                self._add_heading(doc, stripped[4:], 3)
                i += 1
                continue
            elif stripped.startswith("## "):
                self._add_heading(doc, stripped[3:], 2)
                i += 1
                continue
            elif stripped.startswith("# "):
                self._add_heading(doc, stripped[2:], 1)
                i += 1
                continue

            # Horizontal rule
            if re.match(r"^-{3,}$", stripped) or re.match(r"^={3,}$", stripped):
                doc.add_paragraph()
                i += 1
                continue

            # Everything else - formatted paragraph
            self._add_formatted_paragraph(doc, stripped)
            i += 1

    def _process_code_fence(self, doc, lines: List[str], start_index: int) -> int:
        """Process code fence block and return next index."""
        i = start_index + 1
        while i < len(lines) and not lines[i].strip().startswith("```"):
            inner = lines[i].strip()
            if inner:
                # Render content inside fence
                if inner.startswith("#### "):
                    self._add_heading(doc, inner[5:], 4)
                elif inner.startswith("### "):
                    self._add_heading(doc, inner[4:], 3)
                elif inner.startswith("## "):
                    self._add_heading(doc, inner[3:], 2)
                elif inner.startswith("# "):
                    self._add_heading(doc, inner[2:], 1)
                elif inner == "---":
                    doc.add_paragraph()
                else:
                    self._add_formatted_paragraph(doc, inner)
            i += 1

        if i < len(lines):
            i += 1  # Skip closing fence
        return i

    def _process_table(self, doc, lines: List[str], start_index: int) -> int:
        """Process markdown table and return next index."""
        table_lines: List[str] = []
        i = start_index

        while i < len(lines) and lines[i].strip().startswith("|"):
            table_lines.append(lines[i])
            i += 1

        self._insert_professional_table(doc, table_lines)
        return i

    # ===================================================================
    # Heading Helper
    # ===================================================================

    def _add_heading(self, doc, text: str, level: int) -> None:
        """Add a professionally styled heading."""
        clean_text = self._strip_markdown_formatting(text)
        h = doc.add_heading(clean_text, level=level)

        # Headings automatically use the styled heading formats we configured
        # No additional formatting needed

    def _strip_markdown_formatting(self, text: str) -> str:
        """Strip markdown bold/italic/code syntax."""
        text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)  # Bold
        text = re.sub(r'\*(.+?)\*', r'\1', text)      # Italic
        text = re.sub(r'`(.+?)`', r'\1', text)        # Code
        return text

    # ===================================================================
    # Image Insertion
    # ===================================================================

    def _insert_image(self, doc, image_id: str, caption: str, image_lookup: Dict[str, Dict]) -> None:
        """Insert an image with professional caption."""
        Inches = self._Inches
        Pt = self._Pt
        WD_ALIGN = self._WD_ALIGN

        entry = image_lookup.get(image_id, {})
        raw_path = entry.get("storage_path", "")
        display_caption = caption or entry.get("caption", image_id)

        resolved: Optional[Path] = None
        if raw_path:
            p = Path(raw_path)
            if p.exists():
                resolved = p
            else:
                p2 = Path.cwd() / raw_path
                if p2.exists():
                    resolved = p2

        if resolved:
            try:
                # Insert image
                doc.add_picture(str(resolved), width=Inches(self.IMAGE_WIDTH_INCHES))

                # Center the image
                last_para = doc.paragraphs[-1]
                last_para.alignment = WD_ALIGN.CENTER
                last_para.paragraph_format.space_before = Pt(6)
                last_para.paragraph_format.space_after = Pt(3)

                # Add caption
                cap_para = doc.add_paragraph()
                cap_para.alignment = WD_ALIGN.CENTER
                cap_run = cap_para.add_run(display_caption)
                cap_run.font.name = self.FONT_FAMILY
                cap_run.font.size = Pt(self.CAPTION_FONT_SIZE)
                cap_run.font.italic = True
                cap_run.font.color.rgb = self._RGBColor(0x60, 0x60, 0x60)
                cap_para.paragraph_format.space_before = Pt(3)
                cap_para.paragraph_format.space_after = Pt(12)

                return
            except Exception as exc:
                logger.warning(f"Could not embed image '{image_id}': {exc}")

        # Placeholder when image unavailable
        ph_para = doc.add_paragraph()
        ph_para.alignment = WD_ALIGN.CENTER
        ph_run = ph_para.add_run(f"[Figure: {display_caption}]")
        ph_run.font.name = self.FONT_FAMILY
        ph_run.font.size = Pt(self.CAPTION_FONT_SIZE)
        ph_run.font.italic = True
        ph_run.font.color.rgb = self._RGBColor(0xAA, 0xAA, 0xAA)

    # ===================================================================
    # Professional Table Formatting
    # ===================================================================

    def _insert_professional_table(self, doc, table_lines: List[str]) -> None:
        """Convert markdown table to professionally formatted Word table."""
        separator_re = re.compile(r"^\s*\|[\s\-:|]+\|\s*$")

        def parse_row(raw_line: str) -> List[str]:
            return [c.strip() for c in raw_line.strip().strip("|").split("|")]

        data_lines = [l for l in table_lines if not separator_re.match(l)]
        if not data_lines:
            return

        header_cells = parse_row(data_lines[0])
        data_rows = [parse_row(l) for l in data_lines[1:]]
        num_cols = max(len(header_cells), 1)

        # Normalize row lengths
        def pad(row, n):
            return row[:n] + [""] * max(0, n - len(row))

        header_cells = pad(header_cells, num_cols)
        data_rows = [pad(r, num_cols) for r in data_rows]

        try:
            # Create table
            tbl = doc.add_table(rows=1 + len(data_rows), cols=num_cols)
            tbl.style = "Light Grid - Accent 1"  # Professional table style

            # Header row - bold with shading
            hdr_row = tbl.rows[0]
            for col_idx, cell_text in enumerate(header_cells):
                cell = hdr_row.cells[col_idx]
                cell.text = cell_text

                # Format header cell
                for para in cell.paragraphs:
                    for run in para.runs:
                        run.font.name = self.FONT_FAMILY
                        run.font.size = self._Pt(self.BODY_FONT_SIZE)
                        run.font.bold = True
                        run.font.color.rgb = self._RGBColor(0xFF, 0xFF, 0xFF)  # White text
                    para.alignment = self._WD_ALIGN.LEFT

                # Blue header background
                self._shade_cell(cell, "2E75B5")

            # Data rows
            for row_idx, row_cells in enumerate(data_rows):
                row = tbl.rows[row_idx + 1]
                for col_idx, cell_text in enumerate(row_cells):
                    cell = row.cells[col_idx]
                    cell.text = cell_text

                    # Format data cell
                    for para in cell.paragraphs:
                        for run in para.runs:
                            run.font.name = self.FONT_FAMILY
                            run.font.size = self._Pt(self.BODY_FONT_SIZE)
                        para.alignment = self._WD_ALIGN.LEFT

                    # Alternate row shading
                    if row_idx % 2 == 0:
                        self._shade_cell(cell, "F2F2F2")  # Light gray

            # Add spacing after table
            doc.add_paragraph()

        except Exception as exc:
            logger.warning(f"Table insertion failed: {exc}")
            # Fallback to plain text
            for line in data_lines:
                doc.add_paragraph(line.strip("|").strip())

    def _shade_cell(self, cell, fill_hex: str) -> None:
        """Apply background shading to table cell."""
        try:
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn

            tc = cell._tc
            tcPr = tc.get_or_add_tcPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), fill_hex)
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            tcPr.append(shd)
        except Exception:
            pass

    # ===================================================================
    # Formatted Paragraph (with bullets, numbering, inline formatting)
    # ===================================================================

    def _add_formatted_paragraph(self, doc, text: str) -> None:
        """Add paragraph with proper bullets, numbering, and inline formatting."""
        Inches = self._Inches
        Pt = self._Pt

        # Bullet list
        if re.match(r'^[-*•]\s', text):
            item_text = text[1:].lstrip()
            para = doc.add_paragraph(style="List Bullet")
            self._add_inline_runs(para, item_text)
            return

        # Numbered list - render number explicitly
        num_match = re.match(r'^(\d+)\.\s+(.+)$', text)
        if num_match:
            num = num_match.group(1)
            item_text = num_match.group(2)
            para = doc.add_paragraph()
            para.paragraph_format.left_indent = Inches(0.35)
            para.paragraph_format.first_line_indent = Inches(-0.35)
            para.add_run(f"{num}. ")
            self._add_inline_runs(para, item_text)
            return

        # Regular paragraph
        para = doc.add_paragraph()
        self._add_inline_runs(para, text)

    def _add_inline_runs(self, para, text: str) -> None:
        """Parse **bold** and *italic* and add formatted runs."""
        token_re = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*)")

        for part in token_re.split(text):
            if part.startswith("**") and part.endswith("**") and len(part) > 4:
                run = para.add_run(part[2:-2])
                run.bold = True
            elif part.startswith("*") and part.endswith("*") and len(part) > 2:
                run = para.add_run(part[1:-1])
                run.italic = True
            elif part:
                para.add_run(part)

    # ===================================================================
    # Utilities
    # ===================================================================

    def _hex_to_rgb(self, hex_str: str):
        """Convert hex color to RGBColor."""
        h = hex_str.lstrip("#")
        return self._RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))

    @staticmethod
    def _section_sort_key(section_number: str) -> Tuple:
        """Convert section number to sortable tuple."""
        try:
            return tuple(int(x) for x in str(section_number).split("."))
        except (ValueError, AttributeError):
            return (999,)
